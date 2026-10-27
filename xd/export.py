"""Export de l'argumentaire au format Word (.docx)."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .generateur import NIVEAUX, Argumentaire, _cite

VERT = RGBColor(0x1F, 0x6B, 0x52)
GRIS = RGBColor(0x66, 0x66, 0x66)


def vers_docx(arg: Argumentaire, base, afficher_niveaux: bool = True, inclure_notes: bool = False) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    t = doc.add_heading(arg.titre, level=0)
    for r in t.runs:
        r.font.color.rgb = VERT
    p = doc.add_paragraph(arg.sous_titre)
    p.runs[0].italic = True

    for s in arg.sections:
        if s.note_interne and not inclure_notes:
            continue
        h = doc.add_heading(s.titre, level=1)
        for r in h.runs:
            r.font.color.rgb = VERT
        for b in s.blocs:
            if b.type == "paragraphe":
                doc.add_paragraph(b.texte)
            elif b.type == "preuve":
                para = doc.add_paragraph(style="List Bullet")
                if b.titre:
                    para.add_run(b.titre + ". ").bold = True
                para.add_run(b.texte + _cite(b))
                if afficher_niveaux and b.niveau:
                    run = para.add_run(f" ({NIVEAUX.get(b.niveau, b.niveau)})")
                    run.italic = True
                    run.font.color.rgb = GRIS
                if afficher_niveaux and b.limite:
                    note = doc.add_paragraph()
                    note.paragraph_format.left_indent = Pt(24)
                    r = note.add_run("Précaution : " + b.limite)
                    r.italic = True
                    r.font.size = Pt(9)
                    r.font.color.rgb = GRIS
            elif b.type == "puces":
                if b.titre:
                    doc.add_paragraph().add_run(b.titre).bold = True
                for i in b.items:
                    doc.add_paragraph(i, style="List Bullet")
            elif b.type == "verbatim":
                q = doc.add_paragraph()
                q.paragraph_format.left_indent = Pt(18)
                r = q.add_run(f"« {b.texte} »")
                r.italic = True
                q.add_run(f" — {b.attribution} [{b.ref}]").font.color.rgb = GRIS
            elif b.type == "question":
                doc.add_paragraph().add_run(f"« {b.titre} »").bold = True
                doc.add_paragraph(b.texte)
            elif b.type == "tableau" and b.tableau is not None and not b.tableau.empty:
                if b.titre:
                    doc.add_paragraph().add_run(b.titre).bold = True
                df = b.tableau.reset_index() if b.tableau.index.name else b.tableau
                table = doc.add_table(rows=1, cols=len(df.columns))
                table.style = "Light Grid Accent 1"
                for j, c in enumerate(df.columns):
                    table.rows[0].cells[j].text = str(c)
                for _, row in df.iterrows():
                    cells = table.add_row().cells
                    for j, v in enumerate(row):
                        cells[j].text = "" if v is None else str(v)
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            for r in para.runs:
                                r.font.size = Pt(8.5)

    h = doc.add_heading("Sources", level=1)
    for r in h.runs:
        r.font.color.rgb = VERT
    for num, sid in arg.references:
        src = base.sources.get(sid, {})
        statut = "" if src.get("diffusion") == "publique" else " — document interne"
        url = f" — {src['url']}" if src.get("url") else ""
        para = doc.add_paragraph(f"[{num}] {src.get('auteur', '')}, {src.get('titre', sid)}, {src.get('date', '')}{url}{statut}")
        para.runs[0].font.size = Pt(9)
    if afficher_niveaux:
        n = doc.add_paragraph("Niveaux de preuve : donnée mesurée, retour de terrain, texte en vigueur, estimation, "
                              "projection, objectif du programme. Estimations et projections reposent sur les hypothèses "
                              "des documents cités.")
        n.runs[0].italic = True
        n.runs[0].font.size = Pt(8.5)
        n.alignment = WD_ALIGN_PARAGRAPH.LEFT

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
